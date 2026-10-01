"""Auditable download, validation, country matching and scenario pipeline."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import urllib.request
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

WAW_URL = "https://datacatalogfiles.worldbank.org/ddh-published/0039597/DR0095901/What_a_Waste_3.0_COUNTRY_Dataset_%26_Codebook.xlsx"
API = "https://api.worldbank.org/v2/country/all/indicator/{code}?format=json&per_page=20000&date=2015:2024"
INDICATORS = {"gdp_pc_ppp": "NY.GDP.PCAP.PP.KD", "urban_pct": "SP.URB.TOTL.IN.ZS"}
GEN = "msw_total_msw_generated_tonnes_year"
YEAR = "msw_total_msw_generation_year"
COMP = ["composition_msw_" + x + "_percent" for x in (
    "food_organic_waste", "glass", "metal", "paper_cardboard", "plastic", "rubber_leather",
    "wood", "yard_garden_green_waste", "textile_waste", "weee", "hazardous",
    "diapers_and_hygiene", "other")]
ELIGIBLE = ["composition_msw_" + x + "_percent" for x in (
    "food_organic_waste", "glass", "metal", "paper_cardboard", "plastic", "yard_garden_green_waste")]
TREAT = ["waste_treatment_" + x + "_percent" for x in (
    "open_dumpsite", "controlled_landfill", "sanitary_landfill_landfill_gas_system",
    "landfill_unspecified", "anaerobic_digestion", "compost", "recycling", "incineration",
    "mbt", "rdf", "other", "unaccounted_for")]
UNC = "waste_uncollected_percent"
COVER = "waste_collection_coverage_total_percent_of_waste"
RECOVERY = ["waste_treatment_" + x + "_percent" for x in ("anaerobic_digestion", "compost", "recycling")]


def download(url: str, path: Path, kind: str) -> dict:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        req = urllib.request.Request(url, headers={"User-Agent": "buried-resources-research/0.1"})
        with urllib.request.urlopen(req, timeout=90) as response:
            data = response.read(35_000_001)
            if len(data) > 35_000_000:
                raise ValueError(f"Oversize download: {url}")
        path.write_bytes(data)
    b = path.read_bytes()
    if kind == "xlsx" and not b.startswith(b"PK\x03\x04"):
        raise ValueError(f"Expected XLSX ZIP bytes: {url}")
    if kind == "json" and not b.lstrip().startswith(b"["):
        raise ValueError(f"Expected JSON list: {url}")
    return {"url": url, "bytes": len(b), "sha256": hashlib.sha256(b).hexdigest(), "file": path.name}


def read_waw(path: Path) -> pd.DataFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    if "Country dataset" not in wb.sheetnames or "Codebook" not in wb.sheetnames:
        raise ValueError("Missing country data sheet or codebook")
    rows = wb["Country dataset"].values
    next(rows)  # descriptive labels
    keys = [str(x).strip() if x is not None else "" for x in next(rows)]
    required = {"iso3c", "country_name", "region_id", GEN, YEAR, UNC, COVER, *COMP, *TREAT}
    if not required.issubset(keys) or len(keys) != len(set(keys)):
        raise ValueError(f"Missing or repeated columns: {sorted(required - set(keys))}")
    frame = pd.DataFrame(rows, columns=keys)
    frame = frame[frame.iso3c.notna()].copy()
    if frame.iso3c.duplicated().any() or not frame.iso3c.astype(str).str.fullmatch(r"[A-Z]{3}").all():
        raise ValueError("Nonunique or malformed ISO3 codes")
    for col in [GEN, YEAR, UNC, COVER, *COMP, *TREAT]:
        frame[col] = pd.to_numeric(frame[col], errors="coerce")
    return frame


def read_wdi(path: Path, indicator: str) -> pd.DataFrame:
    obj = json.loads(path.read_text())
    if not isinstance(obj, list) or len(obj) != 2 or obj[0].get("pages") != 1 or not isinstance(obj[1], list):
        raise ValueError(f"Invalid or paginated WDI response: {indicator}")
    rows = []
    for item in obj[1]:
        if item["indicator"]["id"] != indicator:
            raise ValueError("Wrong WDI indicator")
        if item["value"] is not None and len(item["countryiso3code"]) == 3:
            rows.append((item["countryiso3code"], int(item["date"]), float(item["value"])))
    if not rows:
        raise ValueError("Empty WDI indicator")
    return pd.DataFrame(rows, columns=["iso3c", "year", "value"])


def validate_and_prepare(raw: pd.DataFrame, wdi: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    x = raw.copy()
    diagnostics = {}
    def gate(name, valid):
        nonlocal x
        valid = valid.fillna(False)
        diagnostics[name] = {"before": len(x), "retained": int(valid.sum()), "excluded": int((~valid).sum())}
        x = x.loc[valid].copy()

    gate("positive_generation_and_valid_year", (x[GEN] > 0) & x[YEAR].between(2010, 2024))
    # Percent variables are decimal fractions in the actual spreadsheet.
    values = x[COMP + TREAT + [UNC, COVER]]
    gate("fractions_within_zero_one", (values.isna() | ((values >= 0) & (values <= 1.000001))).all(axis=1))
    comp_total = x[COMP].sum(axis=1, min_count=1)
    gate("composition_closure_95_105pct", comp_total.between(.95, 1.05) & x[ELIGIBLE].notna().any(axis=1))
    x["eligible_share"] = x[ELIGIBLE].fillna(0).sum(axis=1)
    gate("plausible_eligible_fraction", x.eligible_share.between(.01, 1))
    treat_cols = TREAT + [UNC]
    treat_total = x[treat_cols].sum(axis=1, min_count=1)
    gate("treatment_closure_95_105pct", treat_total.between(.95, 1.05))
    x["uncollected_share"] = x[UNC]
    fallback = x.uncollected_share.isna() & x[COVER].notna()
    x.loc[fallback, "uncollected_share"] = 1 - x.loc[fallback, COVER]
    x["coverage_source"] = np.where(x[UNC].notna(), "uncollected_mass_share", np.where(fallback, "collected_mass_share", "missing"))
    gate("mass_based_collection_coverage", x.uncollected_share.between(0, 1))
    x["collection_share"] = 1 - x.uncollected_share
    x["recovery_share"] = x[RECOVERY].fillna(0).sum(axis=1)
    gate("mass_balance_recovery_feasible", (x.recovery_share <= x.eligible_share * x.collection_share + .01) & (x.recovery_share <= x.collection_share + .01))
    # Closest observed WDI year to the WAW observation year; ties choose earlier.
    for name, table in wdi.items():
        grouped = {iso: g for iso, g in table.groupby("iso3c")}
        selected = []
        years = []
        for iso, year in zip(x.iso3c, x[YEAR]):
            g = grouped.get(iso)
            if g is None:
                selected.append(np.nan); years.append(np.nan); continue
            g = g.assign(distance=(g.year - int(year)).abs()).sort_values(["distance", "year"])
            best = g.iloc[0]
            selected.append(best.value if best.distance <= 3 else np.nan)
            years.append(best.year if best.distance <= 3 else np.nan)
        x[name] = selected
        x[name + "_year"] = years
    gate("wdi_within_three_years", x.gdp_pc_ppp.gt(0) & x.urban_pct.between(0, 100))
    x["capture_given_collection"] = (x.recovery_share / (x.eligible_share * x.collection_share).replace(0, np.nan)).clip(upper=1)
    gate("positive_collection", x.capture_given_collection.notna())
    x["eligible_tonnes"] = x[GEN] * x.eligible_share
    x["recovered_tonnes_proxy"] = x[GEN] * x.recovery_share
    x["unrecovered_eligible_tonnes_proxy"] = (x.eligible_tonnes - x.recovered_tonnes_proxy).clip(lower=0)
    x["collection_gap_proxy"] = x.eligible_tonnes * x.uncollected_share
    x["treatment_gap_proxy"] = (x.eligible_tonnes * x.collection_share - x.recovered_tonnes_proxy).clip(lower=0)
    x["recovery_efficiency_proxy"] = x.recovery_share / x.eligible_share
    cols = ["iso3c", "country_name", "region_id", YEAR, GEN, "gdp_pc_ppp", "gdp_pc_ppp_year", "urban_pct", "urban_pct_year", "eligible_share", "collection_share", "coverage_source", "recovery_share", "capture_given_collection", "recovery_efficiency_proxy", "eligible_tonnes", "recovered_tonnes_proxy", "unrecovered_eligible_tonnes_proxy", "collection_gap_proxy", "treatment_gap_proxy"]
    metrics = x[cols].rename(columns={GEN: "waste_tonnes_year", YEAR: "waste_year"}).reset_index(drop=True)
    return x, metrics, diagnostics


def peer_scenarios(metrics: pd.DataFrame, k: int = 8, minimum: int = 5) -> tuple[pd.DataFrame, pd.DataFrame]:
    # Match without conditioning on outcome. Robust scale prevents a single feature dominating.
    z = np.column_stack([np.log(metrics.gdp_pc_ppp), metrics.urban_pct / 100, metrics.eligible_share])
    center = np.median(z, axis=0)
    scale = np.subtract(*np.percentile(z, [75, 25], axis=0))
    scale = np.where(scale > 1e-8, scale, 1)
    z = (z - center) / scale
    peers = []
    scenarios = []
    for i, target in metrics.iterrows():
        d = np.linalg.norm(z - z[i], axis=1)
        d[i] = np.inf
        nearest = np.argsort(d)[:k]
        if len(nearest) < minimum or not np.isfinite(d[nearest[-1]]):
            continue
        pool = metrics.iloc[nearest]
        for rank, j in enumerate(nearest, start=1):
            peers.append({"iso3c": target.iso3c, "peer_iso3c": metrics.iloc[j].iso3c, "rank": rank, "scaled_distance": float(d[j])})
        c_star = max(target.collection_share, float(pool.collection_share.quantile(.75)))
        t_star = max(target.capture_given_collection, float(pool.capture_given_collection.quantile(.75)))
        # Sequential attribution: expand collection at current capture, then improve capture.
        collection_gain = target.eligible_tonnes * (c_star - target.collection_share) * target.capture_given_collection
        treatment_gain = target.eligible_tonnes * c_star * (t_star - target.capture_given_collection)
        reduction = collection_gain + treatment_gain
        scenarios.append({"iso3c": target.iso3c, "country_name": target.country_name,
            "peer_count": len(nearest), "peer_p75_collection": float(pool.collection_share.quantile(.75)),
            "peer_p75_capture": float(pool.capture_given_collection.quantile(.75)),
            "target_collection": c_star, "target_capture": t_star,
            "collection_gain_tonnes_proxy": collection_gain, "treatment_gain_tonnes_proxy": treatment_gain,
            "additional_recovery_tonnes_proxy": reduction,
            "remaining_unrecovered_tonnes_proxy": max(0, target.unrecovered_eligible_tonnes_proxy - reduction),
            "priority": "no_measured_peer_gap" if reduction <= 1e-9 else
                        ("collection" if collection_gain > treatment_gain else "sorting_and_biological_treatment")})
    return pd.DataFrame(peers), pd.DataFrame(scenarios)


def charts(metrics: pd.DataFrame, scenarios: pd.DataFrame, out: Path):
    plt.rcParams.update({"font.size": 10, "figure.dpi": 140})
    fig, ax = plt.subplots(figsize=(7.2, 5))
    ax.scatter(metrics.eligible_share * 100, metrics.recovery_efficiency_proxy * 100,
               c=metrics.collection_share, cmap="viridis", alpha=.75, s=25)
    ax.set(xlabel="Potentially recoverable composition (%)", ylabel="Observed recovery / eligible composition (%)",
           title="Country-level recovery and composition")
    fig.colorbar(ax.collections[0], ax=ax, label="Collected mass share")
    fig.tight_layout(); fig.savefig(out / "composition_efficiency.png"); plt.close(fig)
    top = scenarios.nlargest(15, "additional_recovery_tonnes_proxy").sort_values("additional_recovery_tonnes_proxy")
    fig, ax = plt.subplots(figsize=(8, 6))
    ax.barh(top.iso3c, top.collection_gain_tonnes_proxy / 1e6, label="Collection", color="#5594a4")
    ax.barh(top.iso3c, top.treatment_gain_tonnes_proxy / 1e6,
            left=top.collection_gain_tonnes_proxy / 1e6, label="Post-collection", color="#db9b52")
    ax.set(xlabel="Additional recovery scenario (million tonnes/year, proxy)", title="Largest descriptive benchmark gaps")
    ax.legend(); fig.tight_layout(); fig.savefig(out / "scenario_top15.png"); plt.close(fig)
    fig, ax = plt.subplots(figsize=(7, 4))
    ax.hist(metrics.waste_year.astype(int), bins=np.arange(2009.5, 2025.6, 1), color="#5594a4")
    ax.set(xlabel="Waste-generation reference year", ylabel="Countries", title="Temporal distribution of the analyzed sample")
    fig.tight_layout(); fig.savefig(out / "sample_years.png"); plt.close(fig)


def run(root: Path):
    raw_dir = root / "data" / "raw"; out = root / "outputs"
    out.mkdir(parents=True, exist_ok=True)
    audit = {"waw": download(WAW_URL, raw_dir / "waw3_country.xlsx", "xlsx")}
    wdi = {}
    for name, code in INDICATORS.items():
        path = raw_dir / f"wdi_{code}.json"
        audit[name] = download(API.format(code=code), path, "json")
        wdi[name] = read_wdi(path, code)
    (out / "source_audit.json").write_text(json.dumps(audit, indent=2), encoding="utf-8")
    original = read_waw(raw_dir / "waw3_country.xlsx")
    _, metrics, flow = validate_and_prepare(original, wdi)
    if len(metrics) < 6:
        raise ValueError(f"Only {len(metrics)} eligible countries; cannot form peers")
    peers, scenarios = peer_scenarios(metrics)
    if scenarios.empty:
        raise ValueError("No peer scenarios produced")
    pd.DataFrame([{"gate": k, **v} for k, v in flow.items()]).to_csv(out / "sample_flow.csv", index=False)
    metrics.to_csv(out / "country_metrics.csv", index=False)
    peers.to_csv(out / "peer_matches.csv", index=False)
    scenarios.to_csv(out / "scenario.csv", index=False)
    summary = {"source_countries": len(original), "analyzed_countries": len(metrics), "peer_scenarios": len(scenarios),
               "waste_year_min": int(metrics.waste_year.min()), "waste_year_max": int(metrics.waste_year.max()),
               "total_waste_tonnes_in_analyzed_sample": float(metrics.waste_tonnes_year.sum()),
               "total_additional_recovery_tonnes_proxy": float(scenarios.additional_recovery_tonnes_proxy.sum()),
               "share_with_collection_as_larger_scenario_gain": float((scenarios.priority == "collection").mean()),
               "scope": "descriptive conditional sample; mixed observation years; counterfactual proxy, not causal or measured landfill diversion"}
    (out / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding="utf-8")
    (out / "validation.json").write_text(json.dumps({"source_rows": len(original), "gates": flow,
        "compositional_tolerance": [0.95, 1.05], "treatment_tolerance": [0.95, 1.05],
        "wdi_max_year_gap": 3, "percent_storage": "decimal fractions"}, indent=2), encoding="utf-8")
    charts(metrics, scenarios, out)
    print(json.dumps(summary, indent=2, ensure_ascii=False))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path.cwd())
    args = parser.parse_args()
    run(args.root.resolve())


if __name__ == "__main__":
    main()
